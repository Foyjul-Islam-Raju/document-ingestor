"""Structured extraction for CSV, XLSX, and legacy XLS workbooks."""

from __future__ import annotations

import csv
import io
import math
import os
import zipfile
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

from .errors import ExtractionError


IS_VERCEL = os.getenv("VERCEL") == "1"
MAX_TABLE_CELLS = int(os.getenv("MAX_TABLE_CELLS", "20000" if IS_VERCEL else "200000"))
MAX_TABLE_ROWS = int(os.getenv("MAX_TABLE_ROWS", "10000" if IS_VERCEL else "50000"))
MAX_TABLE_COLUMNS = int(os.getenv("MAX_TABLE_COLUMNS", "500"))
MAX_SHEETS = int(os.getenv("MAX_SHEETS", "100"))
MAX_UNCOMPRESSED_XLSX_BYTES = int(os.getenv("MAX_UNCOMPRESSED_XLSX_BYTES", str(500 * 1024 * 1024)))


def extract_csv(path: Path) -> dict[str, Any]:
    raw = path.read_bytes()
    text = None
    encoding_used = None
    for encoding in ("utf-8-sig", "utf-16", "cp1252"):
        try:
            text = raw.decode(encoding)
            encoding_used = encoding
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ExtractionError("unsupported_encoding", "The CSV encoding is not supported (tried UTF-8, UTF-16, and Windows-1252).")
    if "\x00" in text:
        raise ExtractionError("corrupt_csv", "The CSV contains invalid null characters.")

    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",\t;|") if sample.strip() else csv.excel
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text, newline=""), dialect)
    rows, truncated = _bounded_rows(reader)
    return {
        "sheet_count": 1,
        "sheets": [{"name": "CSV", "row_count": len(rows), "column_count": max((len(row) for row in rows), default=0), "rows": rows}],
        "warnings": ["CSV data was truncated at the configured row/cell limit."] if truncated else [],
        "delimiter": getattr(dialect, "delimiter", ","),
        "encoding": "UTF-8" if encoding_used == "utf-8-sig" else encoding_used.upper(),
    }


def extract_xlsx(path: Path) -> dict[str, Any]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ExtractionError("dependency_missing", "XLSX support requires openpyxl.") from exc

    try:
        with zipfile.ZipFile(path) as archive:
            members = archive.infolist()
            if sum(member.file_size for member in members) > MAX_UNCOMPRESSED_XLSX_BYTES:
                raise ExtractionError("workbook_too_large", "The XLSX expands beyond the configured safety limit.")
    except ExtractionError:
        raise
    except (zipfile.BadZipFile, OSError) as exc:
        raise ExtractionError("corrupt_xlsx", "The XLSX workbook is damaged or is not a valid ZIP-based workbook.") from exc

    formula_book = value_book = None
    try:
        formula_book = load_workbook(path, read_only=True, data_only=False)
        value_book = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        if formula_book is not None:
            formula_book.close()
        if value_book is not None:
            value_book.close()
        raise ExtractionError("corrupt_xlsx", "The XLSX workbook is damaged or could not be opened.") from exc

    sheets: list[dict[str, Any]] = []
    warnings: list[str] = []
    remaining_cells = MAX_TABLE_CELLS
    remaining_rows = MAX_TABLE_ROWS
    try:
        names = formula_book.sheetnames
        if len(names) > MAX_SHEETS:
            warnings.append(f"Only the first {MAX_SHEETS} of {len(names)} sheets were extracted.")
            names = names[:MAX_SHEETS]
        for name in names:
            if remaining_cells <= 0 or remaining_rows <= 0:
                warnings.append("The workbook row/cell limit was reached; remaining sheets were not extracted.")
                break
            formula_sheet = formula_book[name]
            value_sheet = value_book[name]
            source_columns = formula_sheet.max_column or 0
            source_rows = formula_sheet.max_row or 0
            max_column = min(max(source_columns, 1), MAX_TABLE_COLUMNS, remaining_cells)
            max_row = min(source_rows, remaining_rows, max(1, remaining_cells // max_column))
            if source_rows > max_row or source_columns > MAX_TABLE_COLUMNS:
                warnings.append(f"Sheet '{name}' was truncated at the configured row/cell limit.")
            formula_rows = formula_sheet.iter_rows(max_row=max_row, max_col=max_column, values_only=False)
            value_rows = value_sheet.iter_rows(max_row=max_row, max_col=max_column, values_only=False)
            rows: list[list[Any]] = []
            for f_row, v_row in zip(formula_rows, value_rows):
                values = [_json_value(v_cell.value if v_cell.value is not None else f_cell.value) for f_cell, v_cell in zip(f_row, v_row)]
                rows.append(_trim_trailing_empty(values))
            _trim_empty_rows(rows)
            remaining_cells -= min(remaining_cells, max_column * max_row)
            remaining_rows -= min(remaining_rows, max_row)
            sheets.append({
                "name": name,
                "row_count": len(rows),
                "column_count": max((len(row) for row in rows), default=0),
                "rows": rows,
            })
    except Exception as exc:
        raise ExtractionError("corrupt_xlsx", f"A worksheet could not be read ({_short_error(exc)}).") from exc
    finally:
        formula_book.close()
        value_book.close()
    return {"sheet_count": len(sheets), "sheets": sheets, "warnings": warnings}


def extract_xls(path: Path) -> dict[str, Any]:
    try:
        import xlrd
    except ImportError as exc:
        raise ExtractionError("dependency_missing", "Legacy XLS support requires xlrd.") from exc

    try:
        workbook = xlrd.open_workbook(path, on_demand=True)
    except Exception as exc:
        raise ExtractionError("corrupt_xls", "The XLS workbook is damaged or could not be opened.") from exc

    sheets: list[dict[str, Any]] = []
    warnings: list[str] = []
    remaining_cells = MAX_TABLE_CELLS
    remaining_rows = MAX_TABLE_ROWS
    try:
        names = workbook.sheet_names()
        if len(names) > MAX_SHEETS:
            warnings.append(f"Only the first {MAX_SHEETS} of {len(names)} sheets were extracted.")
            names = names[:MAX_SHEETS]
        for name in names:
            if remaining_cells <= 0 or remaining_rows <= 0:
                warnings.append("The workbook row/cell limit was reached; remaining sheets were not extracted.")
                break
            sheet = workbook.sheet_by_name(name)
            column_count = min(max(sheet.ncols, 1), MAX_TABLE_COLUMNS, remaining_cells)
            row_limit = min(sheet.nrows, remaining_rows, max(1, remaining_cells // column_count))
            if sheet.nrows > row_limit or sheet.ncols > MAX_TABLE_COLUMNS:
                warnings.append(f"Sheet '{name}' was truncated at the configured row/cell limit.")
            rows: list[list[Any]] = []
            for row_index in range(row_limit):
                row = [_xls_value(sheet.cell(row_index, col_index), xlrd, workbook.datemode) for col_index in range(min(sheet.ncols, column_count))]
                rows.append(_trim_trailing_empty(row))
            _trim_empty_rows(rows)
            remaining_cells -= min(remaining_cells, column_count * row_limit)
            remaining_rows -= min(remaining_rows, row_limit)
            sheets.append({
                "name": name,
                "row_count": len(rows),
                "column_count": max((len(row) for row in rows), default=0),
                "rows": rows,
            })
            workbook.unload_sheet(name)
    finally:
        workbook.release_resources()
    return {"sheet_count": len(sheets), "sheets": sheets, "warnings": warnings}


def _bounded_rows(reader: Iterable[list[str]]) -> tuple[list[list[str]], bool]:
    rows: list[list[str]] = []
    cell_count = 0
    truncated = False
    for row in reader:
        if len(rows) >= MAX_TABLE_ROWS:
            truncated = True
            break
        if len(row) > MAX_TABLE_COLUMNS:
            row = row[:MAX_TABLE_COLUMNS]
            truncated = True
        remaining = MAX_TABLE_CELLS - cell_count
        if len(row) > remaining:
            row = row[:max(0, remaining)]
            truncated = True
        rows.append(row)
        cell_count += len(row)
        if cell_count >= MAX_TABLE_CELLS:
            truncated = True
            break
    return rows, truncated


def _trim_trailing_empty(row: list[Any]) -> list[Any]:
    while row and row[-1] is None:
        row.pop()
    return row


def _trim_empty_rows(rows: list[list[Any]]) -> None:
    while rows and not any(value is not None for value in rows[-1]):
        rows.pop()


def _json_value(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, bytes):
        return value.decode("utf-8", errors="replace")
    if isinstance(value, float) and not math.isfinite(value):
        return str(value)
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def _xls_value(cell: Any, xlrd: Any, datemode: int) -> Any:
    if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
        return None
    if cell.ctype == xlrd.XL_CELL_DATE:
        try:
            return xlrd.xldate_as_datetime(cell.value, datemode).isoformat()
        except (ValueError, OverflowError):
            return cell.value
    if cell.ctype == xlrd.XL_CELL_ERROR:
        return xlrd.error_text_from_code.get(cell.value, f"#ERROR:{cell.value}")
    return _json_value(cell.value)


def _short_error(exc: Exception) -> str:
    message = " ".join(str(exc).split()) or exc.__class__.__name__
    return message[:160]

